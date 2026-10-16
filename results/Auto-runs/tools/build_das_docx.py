"""Build the Data Availability Statement (.docx) requested by Digital Discovery.

The title/author/affiliation block is copied from Auto_Visco.docx so the DAS matches the manuscript.
Usage: python tools/build_das_docx.py  (run from results/Auto-runs)
"""
import re
from pathlib import Path

import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from title import TITLE, set_title  # noqa: E402
SRC = HERE / "Auto_Visco.docx"
OUT = HERE / "AutoVisco_DAS.docx"
URL = "https://github.com/AC-Formulations-SDL5/AutoVisco"
WEB = "https://ac-formulations-sdl5.github.io/AutoVisco/"

d = docx.Document(SRC)
body = d.element.body
children = [c for c in body.iterchildren() if c.tag != qn("w:sectPr")]
front = []
for c in children:
    if c.tag == qn("w:p") and "".join(t.text or "" for t in c.iter(qn("w:t"))).strip() == "Abstract":
        break
    front.append(c)
for c in children:
    body.remove(c)
for c in front:  # drop co-author comment anchors from the copied block
    for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
        for el in list(c.iter(qn(tag))):
            el.getparent().remove(el)
    for el in list(c.iter(qn("w:commentReference"))):
        el.getparent().getparent().remove(el.getparent())
set_title(front[0])
sectPr = body.find(qn("w:sectPr"))


def append(el):
    sectPr.addprevious(el)
    return el


def para(text="", style=None, bold=False, size=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Paragraph; **...** spans are bold."""
    p = d.add_paragraph(style=style)
    for i, chunk in enumerate(re.split(r"\*\*", text) if text else []):
        if chunk:
            r = p.add_run(chunk)
            r.bold = bold or i % 2 == 1
            if size:
                r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    append(p._p)
    return p


def rule(cell_or_tbl_pr, edges, sz=8):
    borders = OxmlElement("w:tblBorders" if cell_or_tbl_pr.tag == qn("w:tblPr") else "w:tcBorders")
    for e in edges:
        b = OxmlElement(f"w:{e}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:color"), "000000")
        borders.append(b)
    cell_or_tbl_pr.append(borders)


# ---- title block
para("Data Availability Statement", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
for el in front:
    append(el)

para("Availability of data and code", style="Heading 2", align=None)
para(f"The data and code supporting the findings of this article are openly available in the public GitHub "
     f"repository AutoVisco at {URL}. A browsable data portal with direct downloads of every dataset, the "
     f"accompanying documentation, and the citation information is provided at {WEB}. Further experimental details "
     f"are provided in the Supplementary Information.")

para("Data", style="Heading 2", align=None)
para("All datasets are provided as plain CSV files, organized by manuscript figure (**Table 1**). They comprise the "
     "raw torque–displacement records from the manual and automated descents; the silicone-oil calibration and "
     "reproducibility runs; torque time series; multi-speed measurements for every material family (silicone oils, "
     "poly(ethylene glycol), Sepineo™ 600, Solagum™ SF 306, Carbopol® 980, and the thermal interface materials SYY, "
     "Sekisui B-41 and HY-500); manual reference viscometry readings; and the processed analysis outputs (model "
     "comparison, fitted parameters, bootstrap and leave-one-out validation, and master-curve analysis). A data "
     "dictionary describing every file, column, and unit is provided in data/README.md. The data are released under "
     "the Creative Commons Attribution 4.0 International License (CC BY 4.0).")

para("**Table 1.** Datasets in the AutoVisco repository and the manuscript results they support.").paragraph_format.keep_with_next = True
rows = [
    ("Repository folder", "Contents", "Manuscript"),
    ("data/01_manual_descent", "Manual and automated descents of three silicone oils", "**Fig. 2B–D**, **Fig. 7K**"),
    ("data/02_hitpoints", "Container-contact heights per cell", "**Fig. 4B**"),
    ("data/03_silicone_calibration", "Calibration runs on silicone standards", "**Fig. 4**, **Fig. 7A–G**"),
    ("data/04_reproducibility", "Repeat runs at 12.5 and 37 kcP", "Section 3.1"),
    ("data/05_time_series", "Torque versus time at each height", "**Fig. 7**"),
    ("data/06_materials", "Multi-speed descents for all material families; rheology summary", "**Fig. 8**, **Fig. 9A**"),
    ("data/07_reference_viscometry", "Manual reference viscometry readings", "Reference rheology"),
    ("data/08_analysis_outputs", "Model selection, parameters, bootstrap, leave-one-out, master curve",
     "**Fig. 4**, **Fig. 7H–J**, **Fig. S2**"),
    ("data/09_simulation", "Lubrication-model torque versus gap", "Section 2.2"),
]
table = d.add_table(rows=len(rows), cols=3)
append(table._tbl)
rule(table._tbl.tblPr, ("top", "bottom"))
widths = (Inches(2.1), Inches(2.9), Inches(1.5))
for ri, row in enumerate(rows):
    for ci, text in enumerate(row):
        cell = table.rows[ri].cells[ci]
        cell.width = widths[ci]
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        for i, chunk in enumerate(re.split(r"\*\*", text)):
            if chunk:
                r = p.add_run(chunk)
                r.font.size = Pt(9.5)
                r.bold = ri == 0 or i % 2 == 1
        p.paragraph_format.keep_with_next = ri < len(rows) - 1  # keep the table on one page
        if ri == 0:
            rule(cell._tc.get_or_add_tcPr(), ("bottom",), sz=6)

para("Code", style="Heading 2", align=None)
para("The repository also provides the platform control software (the Python package viscometry, including the "
     "hardware interfaces, real-time drag feedback and hit-point detection, the physics-constrained rheology "
     "inference pipeline, and the web interface), the ESP32 washing-station firmware, configuration files, a test "
     "suite, the notebook that reproduces the material rheology results (**Fig. 8** and **Fig. 9A**), and the scripts "
     "that regenerate the vector figures. The code is released under the MIT License. The proprietary torquemeter "
     "communication library (DVT_COM.dll, AMETEK Brookfield) is not redistributed and must be obtained from the "
     "manufacturer.")

para("Supplementary Information", style="Heading 2", align=None)
para("More information is provided in the Supplementary Information, including the bill of materials "
     "(**Table S1**), the washing-protocol evaluation (**Fig. S1**), the dimensionless collapse of silicone drag "
     "curves (**Fig. S2**), and the washing-subsystem electronics (**Fig. S3**).")

d.save(OUT)
print("wrote", OUT)
