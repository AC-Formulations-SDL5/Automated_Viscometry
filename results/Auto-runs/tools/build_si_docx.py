"""Build the Supplementary Information .docx from the main manuscript .docx.

The title/author/affiliation block and Table S1 are copied from Auto_Visco.docx so
the SI keeps the manuscript's formatting; everything else is rebuilt from AutoVisco_SI.md content.
Usage: python tools/build_si_docx.py  (run from results/Auto-runs)
"""
import copy
import re
from pathlib import Path

import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from title import TITLE, set_title  # noqa: E402
SRC = HERE / "Auto_Visco.docx"
OUT = HERE / "AutoVisco_SI.docx"

d = docx.Document(SRC)
body = d.element.body
children = [c for c in body.iterchildren() if c.tag != qn("w:sectPr")]

# Front matter: everything before the "Abstract" heading.
front = []
for c in children:
    if c.tag == qn("w:p") and "".join(t.text or "" for t in c.iter(qn("w:t"))).strip() == "Abstract":
        break
    front.append(c)

# Table S1 caption + table.
cap_s1 = tbl_s1 = None
for i, c in enumerate(children):
    if c.tag == qn("w:p") and "".join(t.text or "" for t in c.iter(qn("w:t"))).startswith("Table S1"):
        cap_s1, tbl_s1 = c, children[i + 1]
        break
assert front and cap_s1 is not None and tbl_s1.tag == qn("w:tbl")

for c in children:
    body.remove(c)
set_title(front[0])
sectPr = body.find(qn("w:sectPr"))


def append(el):
    sectPr.addprevious(el)
    return el


def para(text="", style=None, bold=False, size=None, align=None, italic=False):
    """Paragraph; **...** spans are bold (used for figure/table references)."""
    p = d.add_paragraph(style=style)
    for i, chunk in enumerate(re.split(r"\*\*", text) if text else []):
        if chunk:
            r = p.add_run(chunk)
            r.bold, r.italic = (bold or i % 2 == 1), italic
            if size:
                r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    append(p._p)
    return p


def body(text):
    return para(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def caption(label, text):
    """Caption in body-text style with a bold label."""
    return para(f"**{label}** {text}", align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def figure(path, width_in=None, height_in=None):
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
    kw = {"width": Inches(width_in)} if width_in else {"height": Inches(height_in)}
    p.add_run().add_picture(str(HERE / path), **kw)


# ---- title block
para("Supplementary Information", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
for el in front:
    append(el)

body("**Data availability:** All data, control software, firmware, analysis notebooks, and figure scripts supporting this work are openly available in a public GitHub repository at https://github.com/AC-Formulations-SDL5/AutoVisco (data portal: https://ac-formulations-sdl5.github.io/AutoVisco/).")

para("Contents", bold=True)
for line in ("S1. Bill of materials (**Table S1**)",
             "S2. Washing-protocol evaluation (**Fig. S1**)",
             "S3. Dimensionless collapse of silicone drag curves (**Fig. S2**)",
             "S4. Washing-subsystem electronics (**Fig. S3**)"):
    para(line).paragraph_format.left_indent = Inches(0.3)

# ---- S1
para("S1. Bill of Materials", style="Heading 2").paragraph_format.page_break_before = True
body("**Table S1** lists the components of the automated viscometry platform. Apart from the rotational "
     "torquemeter, the platform is built from low-cost commercial and 3D-printed parts.")
caption("Table S1.", "Bill of materials for the automated rheology platform.")
append(tbl_s1)

table = docx.table.Table(tbl_s1, d)
fixes = {  # first-column label -> {column index: new text}
    "Torque Senser": {0: "Torque Sensor", 1: "Torquemeter (Max: 7,187.0 dyne·cm)",
                      4: "14,500.00 new (3,000.00 used); incl. spindle"},
    "Rotational Spindle": {1: "Cone-shaped (3° angle, 12.0 mm radius)"},  # cost cell is merged with the torquemeter row
    "Platform Layout": {2: "McMaster-Carr"},
}
done = set()
for row in table.rows:
    label = row.cells[0].text.strip()
    for ci, text in fixes.get(label, {}).items():
        runs = row.cells[ci].paragraphs[0].runs
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        done.add(label)
assert done == set(fixes), done
for row in table.rows:  # single-spaced cells so the table fits on one page
    for cell in row.cells:
        for cp in cell.paragraphs:
            cp.paragraph_format.line_spacing = 1.0
            cp.paragraph_format.space_after = Pt(2)

# ---- S2
para("S2. Washing-Protocol Evaluation", style="Heading 2")
body("Four spindle-cleaning protocols of increasing mechanical complexity were benchmarked on high-viscosity "
     "samples: a static spindle, synchronous co-rotation, lateral oscillation, and a multi-modal combination of "
     "rotation, oscillation, and brushing. Only the multi-modal protocol left no detectable high-viscosity residue "
     "on the cone, and it was adopted as the platform default.")
figure("manuscript_figures/Figure_S1_washing_protocols.png", width_in=6.0)
caption("Figure S1.", "Comparison of washing protocols for the cone-shaped spindle.")

# ---- S3
para("S3. Dimensionless Collapse of Silicone Drag Curves", style="Heading 2")
body("Normalizing each silicone drag sweep by its own maximum and plotting it against the common gap axis "
     "collapses the traces onto a single master shape. This supports factorizing the drag into a "
     "geometry-governed shape function F(h) and a viscosity-dependent amplitude A(μ), which is the assumption "
     "that allows one silicone calibration to transfer to other fluids (main text, Section 2.3.3).")
figure("figures_rheology/normalized_drag_per_sample.png", height_in=7.6)
caption("Figure S2.", "Dimensionless collapse of silicone drag curves, supporting factorization into geometry "
                      "shape and viscosity amplitude.")

# ---- S4
para("S4. Washing-Subsystem Electronics", style="Heading 2")
body("The washing subsystem is controlled by an ESP32 microcontroller on a custom PCB. PWM-driven L298N motor "
     "drivers actuate six pumps and three agitation motors through a multiplexed, diode-isolated "
     "channel-sharing configuration, and detergent, water, and isopropanol rinse cycles are triggered from a "
     "minimal host command interface.")
figure("manuscript_figures/Figure_S3_washing_PCB.png", width_in=4.5)
caption("Figure S3.", "PCB layout and embedded motor-control architecture of the washing subsystem.")

d.save(OUT)
print("wrote", OUT)
